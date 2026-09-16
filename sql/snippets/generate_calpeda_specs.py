"""Builds the Calpeda specs load (SQL) from 'Calpeda Specs (3_17_26).xlsx'.

The workbook has one sheet per spec, each listing the same 2,440 Calpeda pumps by DB_ID and
Item Model - the same models, in the same order, that load_calpeda_items_2026_09_12.sql loaded
from 'Calpeda Data Entry (3_17_26).xlsx'. The SQL this writes matches each row to its item by
brand CLP + exact item_model (DB_ID is only a row number in the workbook, not a database id)
and fills in what Item Entry shows for a pump:

  V & FLA          -> tbl_setup_item_specs.volt_1 / volt_2 / fla_1 / fla_2 (template 'PUMP')
  IMPELLER         -> tbl_setup_item_specs.impeller_id (tbl_setup_item_material, by code)
  HP ... MAX. HEAD -> tbl_setup_item_specs_template title/value rows, in Item Entry's own order
  Connection Type  -> tbl_setup_item_additional_specs.connection_type
  Is Special       -> tbl_setup_item_inventory.is_special_item

Every sheet is checked before anything is written: same DB_IDs and model text on every sheet,
the expected column headers, and only the values Item Entry can show (phase 1/3, FLANGED /
THREADED, TRUE/FALSE).

Normally run through load_calpeda_specs.ps1, which builds the load, runs it and clears the API
cache in one go. On its own:

  python generate_calpeda_specs.py "<specs xlsx>" [output.sql] [--commit]
                                   [--add-missing-items "<Calpeda Data Entry xlsx>"]

Without --commit the SQL is a trial run that rolls itself back. --add-missing-items first adds,
in the same transaction, every pump that has no Calpeda item yet (from the Data Entry workbook)
and every impeller code missing from Material setup - for a database that never had the Calpeda
item load. Nothing that already exists is changed or removed.
"""

import datetime
import os
import sys

import openpyxl

# Item Entry's PUMP template (smpc_inventory_app Data/ENUM_ITEM_SPECS.cs PUMP()), in its order.
# "HORESEPOWER" is misspelt there, and the grid only fills rows whose title matches exactly, so
# the workbook's "HORSEPOWER" is written under the app's spelling.
PUMP_TITLES = [
    ("HP", "HORESEPOWER"),
    ("KW", "KILOWATT"),
    ("SUCTION SIZE", "SUCTION SIZE"),
    ("DISCHARGE SIZE", "DISCHARGE SIZE"),
    ("PHASE (1 OR 3)", "PHASE (1 OR 3)"),
    ("MIN. CAPACITY", "MIN. CAPACITY"),
    ("MAX. CAPACITY", "MAX. CAPACITY"),
    ("MIN. HEAD", "MIN. HEAD"),
    ("MAX. HEAD", "MAX. HEAD"),
]

# Sheet -> (0-based value column, header expected there).
VALUE_COLUMN = {
    "HP": (6, "Value"),
    "KW": (4, "Value"),
    "SUCTION SIZE": (4, "Value"),
    "DISCHARGE SIZE": (4, "Value"),
    "PHASE (1 OR 3)": (4, "PHASE (1 OR 3)"),
    "MIN. CAPACITY": (4, "Value"),
    "MAX. CAPACITY": (4, "Value"),
    "MIN. HEAD": (4, "Hmin"),
    "MAX. HEAD": (4, "Value"),
    "IMPELLER": (4, "IMPELLER"),
    "Is Special": (3, "IS SPECIAL"),
    "Connection Type": (3, "CONNECTION"),
}
VOLT_FLA_COLUMNS = {"Voltage 1": 5, "Voltage 2": 6, "FLA 1": 7, "FLA 2": 8}

# The impeller codes the workbook uses, with the names they carry in Lightspeed_test_fresh's
# Material setup (as of 2026-09-15). The names are only used by --add-missing-items, for a
# database that does not have the codes yet.
IMPELLER_MATERIALS = [
    ("AISL 304", "STAINLESS STEEL 304"),
    ("AISL 316", "STAINLESS STEEL 316"),
    ("GB", "BRONZE"),
    ("GG", "CAST IRON"),
    ("GO", "BRASS"),
]
IMPELLER_CODES = {code for code, _ in IMPELLER_MATERIALS}
CONNECTION_TYPES = {"FLANGED", "THREADED"}  # Item Entry's cmb_connection_type items (spec 17)
PHASES = {"1", "3"}

# The Data Entry workbook's columns, spelt as the workbook spells them ("Tangibilty").
ITEM_COLUMNS = ["Item Name", "Item Class", "Item Brand", "Unit of Measure", "Trade Type",
                "Tangibilty Type", "ITEM MODEL", "Catalogue Year", "Price", "Is Stop Selling?",
                "Long Description"]


def fail(message):
    sys.exit("ERROR: " + message)


def text(value):
    """A cell as Item Entry would show it: numbers without a trailing .0, blanks as ''."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else format(value, "g")
    return str(value).strip()


def amps(value):
    """FLA to at most 2 decimals. Several workbook FLAs are computed (4.304347826086957)."""
    raw = text(value)
    if raw == "":
        return ""
    try:
        number = float(raw)
    except ValueError:
        fail("FLA is not a number: %r" % raw)
    return ("%.2f" % number).rstrip("0").rstrip(".")


def sql(value):
    return "N'" + value.replace("'", "''") + "'"


def read_sheet(workbook, name):
    if name not in workbook.sheetnames:
        fail("sheet %r is missing" % name)
    rows = list(workbook[name].iter_rows(values_only=True))
    header = rows[0]
    data = [r for r in rows[1:] if r and r[0] not in (None, "")]
    return header, data


def parse_args(argv):
    commit = False
    items_source = None
    positional = []
    i = 0
    while i < len(argv):
        if argv[i] == "--commit":
            commit = True
        elif argv[i] == "--add-missing-items":
            if i + 1 >= len(argv):
                fail("--add-missing-items needs the path of the Calpeda Data Entry workbook")
            items_source = argv[i + 1]
            i += 1
        else:
            positional.append(argv[i])
        i += 1
    if not positional:
        fail(__doc__)
    return positional, commit, items_source


def read_items(path, models):
    """The Data Entry rows for the pumps on the specs workbook, checked."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    header = [text(h) for h in rows[0][:len(ITEM_COLUMNS)]]
    if header != ITEM_COLUMNS:
        fail("Data Entry header should be %r, found %r" % (ITEM_COLUMNS, header))

    items = []
    seen = set()
    for number, r in enumerate(rows[1:], start=2):
        if not r or len(r) < len(ITEM_COLUMNS) or r[6] in (None, ""):
            continue
        model = r[6]
        if model in seen:
            fail("Data Entry row %d: model %r appears twice" % (number, model))
        seen.add(model)
        item = {
            "name_code": text(r[0]), "class_code": text(r[1]), "brand_code": text(r[2]),
            "uom_code": text(r[3]), "trade_type": text(r[4]), "tangibility": text(r[5]).upper(),
            "model": model, "catalogue_year": text(r[7]), "price": text(r[8]) or "0",
            "stop_selling": text(r[9]) or "0", "long_description": text(r[10]),
        }
        for key in ("name_code", "class_code", "brand_code", "uom_code", "trade_type", "tangibility"):
            if not item[key]:
                fail("Data Entry row %d: %s is blank" % (number, key))
        if item["brand_code"] != "CLP":
            fail("Data Entry row %d: brand %r - this load is for Calpeda (CLP) only" % (number, item["brand_code"]))
        try:
            float(item["price"])
        except ValueError:
            fail("Data Entry row %d: price %r is not a number" % (number, item["price"]))
        if item["stop_selling"] not in ("0", "1"):
            fail("Data Entry row %d: Is Stop Selling? %r" % (number, item["stop_selling"]))
        items.append(item)

    only_specs = set(models) - seen
    only_items = seen - set(models)
    if only_specs or only_items:
        fail("the Data Entry and Specs workbooks list different pumps: %d only on Specs, %d only on Data Entry"
             % (len(only_specs), len(only_items)))
    return items


def main():
    positional, commit, items_source = parse_args(sys.argv[1:])
    source = positional[0]
    today = datetime.date.today().strftime("%Y_%m_%d")
    output = positional[1] if len(positional) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "load_calpeda_specs_%s.sql" % today)

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)

    header, base = read_sheet(workbook, "V & FLA")
    for title, col in VOLT_FLA_COLUMNS.items():
        if text(header[col]) != title:
            fail("V & FLA column %d should be %r, found %r" % (col + 1, title, header[col]))

    pumps = {}
    order = []
    for r in base:
        db_id = int(r[0])
        model = r[1]
        if not model or db_id in pumps:
            fail("V & FLA: blank model or repeated DB_ID %s" % db_id)
        pumps[db_id] = {
            "model": model,
            "volt_1": text(r[5]), "volt_2": text(r[6]),
            "fla_1": amps(r[7]), "fla_2": amps(r[8] if len(r) > 8 else None),
        }
        order.append(db_id)

    if len({p["model"] for p in pumps.values()}) != len(pumps):
        fail("the same model appears on more than one row")

    for sheet, (col, expected) in VALUE_COLUMN.items():
        header, data = read_sheet(workbook, sheet)
        if text(header[col]) != expected:
            fail("%s column %d should be %r, found %r" % (sheet, col + 1, expected, header[col]))
        seen = set()
        for r in data:
            db_id = int(r[0])
            if db_id not in pumps:
                fail("%s: DB_ID %s is not on V & FLA" % (sheet, db_id))
            if r[1] != pumps[db_id]["model"]:
                fail("%s: DB_ID %s model %r differs from V & FLA %r" % (sheet, db_id, r[1], pumps[db_id]["model"]))
            seen.add(db_id)
            pumps[db_id][sheet] = text(r[col] if len(r) > col else None)
        if seen != set(pumps):
            fail("%s is missing %d pumps" % (sheet, len(set(pumps) - seen)))

    for db_id in order:
        p = pumps[db_id]
        if p["PHASE (1 OR 3)"] not in PHASES:
            fail("DB_ID %s: phase %r" % (db_id, p["PHASE (1 OR 3)"]))
        if p["IMPELLER"] not in IMPELLER_CODES:
            fail("DB_ID %s: impeller %r" % (db_id, p["IMPELLER"]))
        p["Connection Type"] = p["Connection Type"].upper()
        if p["Connection Type"] not in CONNECTION_TYPES:
            fail("DB_ID %s: connection %r" % (db_id, p["Connection Type"]))
        p["Is Special"] = p["Is Special"].upper()
        if p["Is Special"] not in ("TRUE", "FALSE"):
            fail("DB_ID %s: is special %r" % (db_id, p["Is Special"]))

    items = read_items(items_source, [pumps[i]["model"] for i in order]) if items_source else None

    write_sql(output, os.path.basename(source), order, pumps, commit, items,
              os.path.basename(items_source) if items_source else None)
    print("checked %d pumps%s, wrote %s (%s)" % (
        len(order), " and their %d items" % len(items) if items else "", output,
        "commit" if commit else "trial run"))


def write_sql(output, source_name, order, pumps, commit=False, items=None, items_name=None):
    out = []
    w = out.append
    w("-- Calpeda pump specs from '%s' - generated by generate_calpeda_specs.py." % source_name)
    w("--")
    w("-- Matches each workbook row to its item by brand CLP + exact item_model: the models")
    w("-- load_calpeda_items_2026_09_12.sql loaded from 'Calpeda Data Entry (3_17_26).xlsx'.")
    w("-- Stops without writing anything if any model does not match exactly one item, or if")
    w("-- an impeller code is missing from Material setup.")
    if items:
        w("--")
        w("-- Also adds, first and in the same transaction, every pump with no item yet (from")
        w("-- '%s') and every impeller code missing from Material setup. Nothing" % items_name)
        w("-- that already exists is changed or removed.")
    w("--")
    w("-- Per pump:")
    w("--   tbl_setup_item_specs            template PUMP, volt_1/2, fla_1/2, impeller_id")
    w("--   tbl_setup_item_specs_template   the 9 PUMP spec rows, replaced as a set")
    w("--   tbl_setup_item_additional_specs connection_type")
    w("--   tbl_setup_item_inventory        is_special_item")
    w("-- manufacturer_origin and every other field are left as they are. Safe to run again.")
    w("--")
    w("-- With @commit = 0 everything is done and counted, then rolled back. load_calpeda_specs.ps1")
    w("-- builds this, runs it (kept with -Commit) and then clears the API cache in one go - item")
    w("-- specs are cached in Redis for up to an hour, and Item Entry would keep the old ones.")
    w("")
    w("SET NOCOUNT ON;")
    w("SET XACT_ABORT ON;")
    w("")
    w("DECLARE @commit bit = %d;" % (1 if commit else 0))
    w("")
    w("IF OBJECT_ID('tempdb..#specs') IS NOT NULL DROP TABLE #specs;")
    w("CREATE TABLE #specs (")
    w("    row_no int NOT NULL PRIMARY KEY,")
    w("    item_model nvarchar(400) NOT NULL,")
    w("    volt_1 nvarchar(50) NOT NULL, volt_2 nvarchar(50) NOT NULL,")
    w("    fla_1 nvarchar(50) NOT NULL, fla_2 nvarchar(50) NOT NULL,")
    w("    impeller_code nvarchar(50) NOT NULL,")
    w("    horsepower nvarchar(50) NOT NULL, kilowatt nvarchar(50) NOT NULL,")
    w("    suction_size nvarchar(50) NOT NULL, discharge_size nvarchar(50) NOT NULL,")
    w("    phase nvarchar(10) NOT NULL,")
    w("    min_capacity nvarchar(50) NOT NULL, max_capacity nvarchar(50) NOT NULL,")
    w("    min_head nvarchar(50) NOT NULL, max_head nvarchar(50) NOT NULL,")
    w("    connection_type nvarchar(20) NOT NULL,")
    w("    is_special bit NOT NULL")
    w(");")
    w("")

    columns = ["row_no", "item_model", "volt_1", "volt_2", "fla_1", "fla_2", "impeller_code",
               "horsepower", "kilowatt", "suction_size", "discharge_size", "phase",
               "min_capacity", "max_capacity", "min_head", "max_head", "connection_type", "is_special"]
    chunk = 500  # a VALUES list takes at most 1,000 rows
    for start in range(0, len(order), chunk):
        w("INSERT INTO #specs (%s) VALUES" % ", ".join(columns))
        lines = []
        for db_id in order[start:start + chunk]:
            p = pumps[db_id]
            values = [str(db_id), sql(p["model"]), sql(p["volt_1"]), sql(p["volt_2"]),
                      sql(p["fla_1"]), sql(p["fla_2"]), sql(p["IMPELLER"])]
            values += [sql(p[sheet]) for sheet, _ in PUMP_TITLES]
            values += [sql(p["Connection Type"]), "1" if p["Is Special"] == "TRUE" else "0"]
            lines.append("(" + ", ".join(values) + ")")
        w(",\n".join(lines) + ";")
        w("")

    if items:
        write_items_table(w, items, chunk)

    w("""BEGIN TRANSACTION;

DECLARE @brand_id bigint = (SELECT id FROM tbl_setup_item_brand WHERE code = 'CLP');
IF @brand_id IS NULL
BEGIN
    RAISERROR('Brand CLP is missing.', 16, 1);
    ROLLBACK TRANSACTION; RETURN;
END
""")

    if items:
        write_items_load(w)

    title_rows = ",\n".join(
        "        (%d, N'%s', s.%s)" % (i + 1, title, column)
        for i, ((_, title), column) in enumerate(zip(PUMP_TITLES, [
            "horsepower", "kilowatt", "suction_size", "discharge_size", "phase",
            "min_capacity", "max_capacity", "min_head", "max_head"])))

    w("""-- Every impeller code must already be a Material setup row.
IF EXISTS (SELECT 1 FROM #specs s WHERE NOT EXISTS (SELECT 1 FROM tbl_setup_item_material m WHERE m.code = s.impeller_code))
BEGIN
    SELECT DISTINCT impeller_code AS missing_material_code FROM #specs s
    WHERE NOT EXISTS (SELECT 1 FROM tbl_setup_item_material m WHERE m.code = s.impeller_code);
    RAISERROR('Impeller codes missing from tbl_setup_item_material - see the result above.', 16, 1);
    ROLLBACK TRANSACTION; RETURN;
END

-- Each workbook model must match exactly one Calpeda item.
IF OBJECT_ID('tempdb..#map') IS NOT NULL DROP TABLE #map;
SELECT s.row_no, i.id AS item_id
INTO #map
FROM #specs s
JOIN tbl_setup_item i ON i.item_brand_id = @brand_id AND i.item_model = s.item_model;

DECLARE @rows int = (SELECT COUNT(*) FROM #specs);
DECLARE @matched int = (SELECT COUNT(DISTINCT row_no) FROM #map);
DECLARE @repeated int = (SELECT COUNT(*) FROM (SELECT row_no FROM #map GROUP BY row_no HAVING COUNT(*) > 1) d);
PRINT 'workbook pumps ' + CAST(@rows AS varchar(10)) + ', matched ' + CAST(@matched AS varchar(10))
    + ', matching more than one item ' + CAST(@repeated AS varchar(10));

IF @matched <> @rows OR @repeated > 0
BEGIN
    SELECT TOP 50 s.row_no, s.item_model AS unmatched_model FROM #specs s
    WHERE NOT EXISTS (SELECT 1 FROM #map m WHERE m.row_no = s.row_no) ORDER BY s.row_no;
    RAISERROR('Not every workbook model matches exactly one Calpeda item - nothing was changed.', 16, 1);
    ROLLBACK TRANSACTION; RETURN;
END

-- 1. Item specs: one row per item, updated in place or created.
UPDATE sp
SET sp.template = N'PUMP',
    sp.volt_1 = s.volt_1, sp.volt_2 = s.volt_2,
    sp.fla_1 = s.fla_1, sp.fla_2 = s.fla_2,
    sp.impeller_id = mat.id
FROM tbl_setup_item_specs sp
JOIN #map m ON m.item_id = sp.based_id
JOIN #specs s ON s.row_no = m.row_no
JOIN tbl_setup_item_material mat ON mat.code = s.impeller_code;
DECLARE @specs_updated int = @@ROWCOUNT;

INSERT INTO tbl_setup_item_specs (based_id, template, volt_1, volt_2, fla_1, fla_2, impeller_id, manufacturer_origin)
SELECT m.item_id, N'PUMP', s.volt_1, s.volt_2, s.fla_1, s.fla_2, mat.id, N''
FROM #map m
JOIN #specs s ON s.row_no = m.row_no
JOIN tbl_setup_item_material mat ON mat.code = s.impeller_code
WHERE NOT EXISTS (SELECT 1 FROM tbl_setup_item_specs sp WHERE sp.based_id = m.item_id)
ORDER BY m.row_no;
DECLARE @specs_inserted int = @@ROWCOUNT;

-- 2. The 9 PUMP spec rows, replaced as a set, in Item Entry's order.
DELETE t
FROM tbl_setup_item_specs_template t
JOIN tbl_setup_item_specs sp ON sp.id = t.based_id
JOIN #map m ON m.item_id = sp.based_id;
DECLARE @template_deleted int = @@ROWCOUNT;

INSERT INTO tbl_setup_item_specs_template (based_id, title, value)
SELECT sp.id, v.title, v.value
FROM #map m
JOIN #specs s ON s.row_no = m.row_no
JOIN tbl_setup_item_specs sp ON sp.based_id = m.item_id
CROSS APPLY (VALUES
%s
    ) v (sort, title, value)
ORDER BY m.row_no, v.sort;
DECLARE @template_inserted int = @@ROWCOUNT;

-- 3. Connection type.
UPDATE a
SET a.connection_type = s.connection_type
FROM tbl_setup_item_additional_specs a
JOIN #map m ON m.item_id = a.based_id
JOIN #specs s ON s.row_no = m.row_no;
DECLARE @connection_updated int = @@ROWCOUNT;

INSERT INTO tbl_setup_item_additional_specs (based_id, connection_type)
SELECT m.item_id, s.connection_type
FROM #map m
JOIN #specs s ON s.row_no = m.row_no
WHERE NOT EXISTS (SELECT 1 FROM tbl_setup_item_additional_specs a WHERE a.based_id = m.item_id)
ORDER BY m.row_no;
DECLARE @connection_inserted int = @@ROWCOUNT;

-- 4. Special item flag.
UPDATE x
SET x.is_special_item = s.is_special
FROM tbl_setup_item_inventory x
JOIN #map m ON m.item_id = x.based_id
JOIN #specs s ON s.row_no = m.row_no;
DECLARE @inventory_updated int = @@ROWCOUNT;

INSERT INTO tbl_setup_item_inventory (based_id, is_special_item)
SELECT m.item_id, s.is_special
FROM #map m
JOIN #specs s ON s.row_no = m.row_no
WHERE NOT EXISTS (SELECT 1 FROM tbl_setup_item_inventory x WHERE x.based_id = m.item_id)
ORDER BY m.row_no;
DECLARE @inventory_inserted int = @@ROWCOUNT;

PRINT 'item specs:       updated ' + CAST(@specs_updated AS varchar(10)) + ', inserted ' + CAST(@specs_inserted AS varchar(10));
PRINT 'spec rows:        deleted ' + CAST(@template_deleted AS varchar(10)) + ', inserted ' + CAST(@template_inserted AS varchar(10));
PRINT 'connection type:  updated ' + CAST(@connection_updated AS varchar(10)) + ', inserted ' + CAST(@connection_inserted AS varchar(10));
PRINT 'special item:     updated ' + CAST(@inventory_updated AS varchar(10)) + ', inserted ' + CAST(@inventory_inserted AS varchar(10));

-- Spot check: the first pump as Item Entry will read it.
SELECT TOP 1 i.item_code, i.item_model, sp.template, sp.volt_1, sp.volt_2, sp.fla_1, sp.fla_2,
    mat.code AS impeller, a.connection_type, x.is_special_item
FROM #map m
JOIN tbl_setup_item i ON i.id = m.item_id
JOIN tbl_setup_item_specs sp ON sp.based_id = i.id
LEFT JOIN tbl_setup_item_material mat ON mat.id = sp.impeller_id
LEFT JOIN tbl_setup_item_additional_specs a ON a.based_id = i.id
LEFT JOIN tbl_setup_item_inventory x ON x.based_id = i.id
ORDER BY m.row_no;

SELECT t.title, t.value
FROM tbl_setup_item_specs_template t
JOIN tbl_setup_item_specs sp ON sp.id = t.based_id
WHERE sp.based_id = (SELECT item_id FROM #map WHERE row_no = (SELECT MIN(row_no) FROM #map))
ORDER BY t.id;

IF @commit = 1
BEGIN
    COMMIT TRANSACTION;
    PRINT 'COMMITTED';
END
ELSE
BEGIN
    ROLLBACK TRANSACTION;
    PRINT 'DRY RUN - rolled back. Set @commit = 1 to keep the changes.';
END""" % title_rows)

    with open(output, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write("\n".join(out) + "\n")


def write_items_table(w, items, chunk):
    w("IF OBJECT_ID('tempdb..#items') IS NOT NULL DROP TABLE #items;")
    w("CREATE TABLE #items (")
    w("    row_no int NOT NULL PRIMARY KEY,")
    w("    name_code nvarchar(256) NOT NULL, class_code nvarchar(256) NOT NULL,")
    w("    brand_code nvarchar(256) NOT NULL, uom_code nvarchar(256) NOT NULL,")
    w("    trade_type nvarchar(50) NOT NULL, tangibility nvarchar(50) NOT NULL,")
    w("    item_model nvarchar(400) NOT NULL, catalogue_year nvarchar(20) NOT NULL,")
    w("    price float NOT NULL, is_stop_selling bit NOT NULL,")
    w("    long_description nvarchar(max) NOT NULL")
    w(");")
    w("")
    columns = ("row_no, name_code, class_code, brand_code, uom_code, trade_type, tangibility, "
               "item_model, catalogue_year, price, is_stop_selling, long_description")
    for start in range(0, len(items), chunk):
        w("INSERT INTO #items (%s) VALUES" % columns)
        lines = []
        for number, it in enumerate(items[start:start + chunk], start=start + 1):
            lines.append("(%d, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)" % (
                number, sql(it["name_code"]), sql(it["class_code"]), sql(it["brand_code"]),
                sql(it["uom_code"]), sql(it["trade_type"]), sql(it["tangibility"]), sql(it["model"]),
                sql(it["catalogue_year"]), repr(float(it["price"])), it["stop_selling"],
                sql(it["long_description"])))
        w(",\n".join(lines) + ";")
        w("")


def write_items_load(w):
    materials = ",\n".join("    (%s, %s)" % (sql(code), sql(name)) for code, name in IMPELLER_MATERIALS)
    w("""-- 0. Pumps with no item yet, and impeller codes missing from Material setup (--add-missing-items).
--    Additive only: nothing that exists is changed. Item codes continue from the highest issued.
INSERT INTO tbl_setup_item_material (code, name)
SELECT v.code, v.name
FROM (VALUES
%s
) v (code, name)
WHERE NOT EXISTS (SELECT 1 FROM tbl_setup_item_material m WHERE m.code = v.code);
DECLARE @materials_added int = @@ROWCOUNT;

-- Each name, class, brand and unit code must name exactly one setup row, or a pump would be
-- added twice or not at all.
IF EXISTS (
    SELECT 1 FROM (SELECT DISTINCT name_code, class_code, brand_code, uom_code FROM #items) c
    WHERE (SELECT COUNT(*) FROM tbl_setup_item_name WHERE code = c.name_code) <> 1
       OR (SELECT COUNT(*) FROM tbl_setup_item_class WHERE code = c.class_code) <> 1
       OR (SELECT COUNT(*) FROM tbl_setup_item_brand WHERE code = c.brand_code) <> 1
       OR (SELECT COUNT(*) FROM tbl_setup_item_unit_measurement WHERE code = c.uom_code) <> 1)
BEGIN
    RAISERROR('An item name, class, brand or unit code on the Data Entry workbook is missing from setup, or set up twice.', 16, 1);
    ROLLBACK TRANSACTION; RETURN;
END

-- Calpeda model names are unique across the whole catalogue (spec 4.2.2): stop if another
-- brand already uses one.
IF EXISTS (SELECT 1 FROM #items i JOIN tbl_setup_item t ON t.item_model = i.item_model
           JOIN tbl_setup_item_brand b ON b.id = t.item_brand_id WHERE b.code <> i.brand_code)
BEGIN
    SELECT TOP 50 i.item_model AS model_used_by_another_brand, b.code AS brand
    FROM #items i JOIN tbl_setup_item t ON t.item_model = i.item_model
    JOIN tbl_setup_item_brand b ON b.id = t.item_brand_id WHERE b.code <> i.brand_code;
    RAISERROR('Workbook models already used by another brand - see the result above.', 16, 1);
    ROLLBACK TRANSACTION; RETURN;
END

DECLARE @last_code int = ISNULL((SELECT MAX(TRY_CAST(item_code AS int)) FROM tbl_setup_item), 0);
DECLARE @items_missing int = (SELECT COUNT(*) FROM #items i WHERE NOT EXISTS (
    SELECT 1 FROM tbl_setup_item t JOIN tbl_setup_item_brand tb ON tb.id = t.item_brand_id
    WHERE tb.code = i.brand_code AND t.item_model = i.item_model));

IF OBJECT_ID('tempdb..#new_items') IS NOT NULL DROP TABLE #new_items;
CREATE TABLE #new_items (item_id bigint NOT NULL, item_model nvarchar(400) NOT NULL);

WITH missing AS (
    SELECT i.*, ROW_NUMBER() OVER (ORDER BY i.row_no) AS seq
    FROM #items i
    WHERE NOT EXISTS (SELECT 1 FROM tbl_setup_item t JOIN tbl_setup_item_brand tb ON tb.id = t.item_brand_id
                      WHERE tb.code = i.brand_code AND t.item_model = i.item_model)
)
INSERT INTO tbl_setup_item (item_name_id, item_class_id, item_brand_id, unit_of_measure_id, item_model,
    catalogue_year, item_code, item_tangibility_type, is_stop_selling, price)
OUTPUT inserted.id, inserted.item_model INTO #new_items (item_id, item_model)
SELECT nm.id, cl.id, br.id, um.id, m.item_model, m.catalogue_year,
    CASE WHEN @last_code + m.seq < 1000 THEN RIGHT('0000' + CAST(@last_code + m.seq AS varchar(10)), 4)
         ELSE CAST(@last_code + m.seq AS varchar(10)) END,
    m.tangibility, m.is_stop_selling, m.price
FROM missing m
JOIN tbl_setup_item_name nm ON nm.code = m.name_code
JOIN tbl_setup_item_class cl ON cl.code = m.class_code
JOIN tbl_setup_item_brand br ON br.code = m.brand_code
JOIN tbl_setup_item_unit_measurement um ON um.code = m.uom_code
ORDER BY m.seq;
DECLARE @items_added int = @@ROWCOUNT;

IF @items_added <> @items_missing
BEGIN
    RAISERROR('Adding the missing pumps did not add exactly one item each - nothing was changed.', 16, 1);
    ROLLBACK TRANSACTION; RETURN;
END

INSERT INTO tbl_setup_item_additional_specs (based_id, long_description)
SELECT n.item_id, i.long_description
FROM #new_items n JOIN #items i ON i.item_model = n.item_model
ORDER BY i.row_no;

INSERT INTO tbl_item_trade_type (based_id, value)
SELECT n.item_id, i.trade_type
FROM #new_items n JOIN #items i ON i.item_model = n.item_model
ORDER BY i.row_no;

PRINT 'impeller codes added: ' + CAST(@materials_added AS varchar(10)) + ', pumps added as items: ' + CAST(@items_added AS varchar(10));
""" % materials)


if __name__ == "__main__":
    main()
